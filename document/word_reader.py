import docx

def extract_text_from_docx(file_path: str) -> str:
    try:
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    except Exception as e:
        raise Exception(f"Gagal membaca file Word: {str(e)}")

def save_text_to_docx(text: str, file_path: str):
    try:
        doc = docx.Document()
        doc.add_paragraph(text)
        doc.save(file_path)
    except Exception as e:
        raise Exception(f"Gagal menyimpan file Word: {str(e)}")

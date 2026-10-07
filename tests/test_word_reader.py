import unittest
import os
import docx
from document.word_reader import extract_text_from_docx

class TestWordReader(unittest.TestCase):
    def setUp(self):
        self.test_file = "test_doc.docx"
        doc = docx.Document()
        doc.add_paragraph("Baris 1")
        doc.add_paragraph("Baris 2")
        doc.save(self.test_file)

    def tearDown(self):
        if os.path.exists(self.test_file):
            os.remove(self.test_file)

    def test_extract_text(self):
        text = extract_text_from_docx(self.test_file)
        self.assertIn("Baris 1", text)
        self.assertIn("Baris 2", text)

if __name__ == '__main__':
    unittest.main()
